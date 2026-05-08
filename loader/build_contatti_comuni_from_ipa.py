#!/usr/bin/env python3
"""
Scarica open data IPA (AgID) e genera un CSV compatibile con load_contatti_comuni.sh:
  pro_com,sito_web,email,pec,telefono,codice_fiscale,indirizzo_fisico

Fonti:
- Dataset "Enti" (enti.xlsx): sito, CF, indirizzo, mail/pec, codice_comune_ISTAT (6 cifre)
- Dataset "Unità Organizzative" (unita-organizzative.xlsx): telefono (per comune ISTAT)

Note:
- IPA usa Codice_comune_ISTAT a 6 cifre (codice ISTAT del comune sede) → mappiamo su pro_com (6 cifre).
- Il telefono non è sul dataset Enti; lo stimiamo scegliendo un numero da una UO del medesimo comune,
  preferendo descrizioni tipo URP/Protocollo/Segreteria se presenti.
"""

from __future__ import annotations

import csv
import re
import sys
import urllib.request
from pathlib import Path


ENTI_XLSX_URL = (
    "https://indicepa.gov.it/ipa-dati/dataset/"
    "5baa3eb8-266e-455a-8de8-b1f434c279b2/resource/"
    "d09adf99-dc10-4349-8c53-27b1e5aa97b6/download/enti.xlsx"
)
UO_XLSX_URL = (
    "https://indicepa.gov.it/ipa-dati/dataset/"
    "c8d2e2b3-a9f1-4123-bc8b-26315ed20fce/resource/"
    "b0aa1f6c-f135-4c8a-b416-396fed4e1a5d/download/unita-organizzative.xlsx"
)


def idx_map(header_row):
    m = {}
    for i, h in enumerate(header_row):
        if h is None:
            continue
        name = str(h).strip()
        if name:
            m[name] = i
    return m


def _btrim(v: object) -> str:
    if v is None:
        return ""
    return str(v).strip()


def _norm_mail(v: str) -> str:
    v = v.strip()
    return v if v else ""


def _pick_first(values: list[str]) -> str:
    for v in values:
        if v:
            return v
    return ""


def _is_pec(mail: str, tipo: str) -> bool:
    t = (tipo or "").strip().lower()
    if t == "pec":
        return True
    # fallback: alcuni dataset legacy possono non valorizzare Tipo_Mail
    return mail.lower().endswith(".pec.it")


def _format_address(indirizzo: str, cap: str) -> str:
    indirizzo = indirizzo.strip()
    cap = cap.strip()
    if not indirizzo and not cap:
        return ""
    if indirizzo and cap:
        return f"{indirizzo}, {cap}"
    return indirizzo or cap


def _download(url: str, out_path: Path) -> None:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    req = urllib.request.Request(
        url,
        headers={"User-Agent": "WEBMAPP municipalities loader (IPA download)"},
        method="GET",
    )
    with urllib.request.urlopen(req, timeout=60) as r, out_path.open("wb") as f:
        f.write(r.read())


def _score_uo(desc: str) -> int:
    d = desc.lower()
    score = 0
    if "urp" in d or "relazioni con il pubblico" in d:
        score += 50
    if "protocollo" in d:
        score += 40
    if "segreter" in d:
        score += 30
    if "anagrafe" in d:
        score += 10
    return score


def main(argv: list[str]) -> int:
    if len(argv) not in (1, 2):
        print("Usage: build_contatti_comuni_from_ipa.py [OUT.csv]", file=sys.stderr)
        return 2

    out_csv = Path(argv[1]) if len(argv) == 2 else None
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("ERROR: python3-openpyxl required", file=sys.stderr)
        return 1

    workdir = Path("/tmp/ipa")
    enti_xlsx = workdir / "enti.xlsx"
    uo_xlsx = workdir / "unita-organizzative.xlsx"

    if not enti_xlsx.exists():
        print(f"Downloading IPA Enti: {ENTI_XLSX_URL}", file=sys.stderr)
        _download(ENTI_XLSX_URL, enti_xlsx)
    if not uo_xlsx.exists():
        print(f"Downloading IPA Unità Organizzative: {UO_XLSX_URL}", file=sys.stderr)
        _download(UO_XLSX_URL, uo_xlsx)

    # 1) Telefoni da UO per Codice_comune_ISTAT
    phone_by_comune: dict[int, str] = {}
    wb_uo = load_workbook(uo_xlsx, read_only=True, data_only=True)
    ws_uo = wb_uo.active
    it_uo = ws_uo.iter_rows(values_only=True)
    hdr_uo = next(it_uo, None)
    if not hdr_uo:
        print("ERROR: unita-organizzative.xlsx empty sheet", file=sys.stderr)
        return 1
    col_uo = idx_map(hdr_uo)
    need_uo = ["Codice_comune_ISTAT", "Descrizione_uo", "Telefono"]
    for k in need_uo:
        if k not in col_uo:
            print(f"ERROR: colonna '{k}' mancante in unita-organizzative.xlsx", file=sys.stderr)
            return 1

    best_score: dict[int, int] = {}
    for row in it_uo:
        if not row:
            continue
        raw_code = _btrim(row[col_uo["Codice_comune_ISTAT"]])
        if not raw_code:
            continue
        if not re.fullmatch(r"\d{6}", raw_code):
            continue
        comune_code = int(raw_code)
        tel = _btrim(row[col_uo["Telefono"]])
        if not tel:
            continue
        desc = _btrim(row[col_uo["Descrizione_uo"]])
        sc = _score_uo(desc)
        prev = best_score.get(comune_code, -1)
        if sc > prev:
            best_score[comune_code] = sc
            phone_by_comune[comune_code] = tel

    wb_uo.close()

    # 2) Record per Comuni dal dataset Enti
    wb_e = load_workbook(enti_xlsx, read_only=True, data_only=True)
    ws_e = wb_e.active
    it_e = ws_e.iter_rows(values_only=True)
    hdr_e = next(it_e, None)
    if not hdr_e:
        print("ERROR: enti.xlsx empty sheet", file=sys.stderr)
        return 1
    col_e = idx_map(hdr_e)

    req = [
        "Tipologia",
        "Codice_comune_ISTAT",
        "Codice_fiscale_ente",
        "Indirizzo",
        "CAP",
        "Sito_istituzionale",
        "Mail1",
        "Tipo_Mail1",
        "Mail2",
        "Tipo_Mail2",
        "Mail3",
        "Tipo_Mail3",
        "Mail4",
        "Tipo_Mail4",
        "Mail5",
        "Tipo_Mail5",
    ]
    for k in req:
        if k not in col_e:
            print(f"ERROR: colonna '{k}' mancante in enti.xlsx", file=sys.stderr)
            return 1

    # IPA: i comuni risultano in genere con Codice_Categoria = "L6" (verificabile su enti.xlsx).
    # Per ogni codice comune ISTAT, prendiamo il primo record.
    seen: set[int] = set()
    rows_out: list[list[object]] = []
    n_total = 0
    n_comuni = 0
    for row in it_e:
        n_total += 1
        if not row:
            continue
        cat = _btrim(row[col_e["Codice_Categoria"]]).upper()
        if cat != "L6":
            continue
        raw_code = _btrim(row[col_e["Codice_comune_ISTAT"]])
        if not re.fullmatch(r"\d{6}", raw_code):
            continue
        pro_com = int(raw_code)
        if pro_com in seen:
            continue
        seen.add(pro_com)
        n_comuni += 1

        sito = _btrim(row[col_e["Sito_istituzionale"]])
        cf = _btrim(row[col_e["Codice_fiscale_ente"]])
        indirizzo = _btrim(row[col_e["Indirizzo"]])
        cap = _btrim(row[col_e["CAP"]])
        indir = _format_address(indirizzo, cap)

        mails = []
        pecs = []
        for i in range(1, 6):
            m = _norm_mail(_btrim(row[col_e[f"Mail{i}"]]))
            t = _btrim(row[col_e[f"Tipo_Mail{i}"]])
            if not m:
                continue
            if _is_pec(m, t):
                pecs.append(m)
            else:
                mails.append(m)

        email = _pick_first(mails)
        pec = _pick_first(pecs)
        tel = phone_by_comune.get(pro_com, "")

        rows_out.append([pro_com, sito, email, pec, tel, cf, indir])

    wb_e.close()

    out_f = out_csv.open("w", newline="", encoding="utf-8") if out_csv else sys.stdout
    close_out = out_csv is not None
    try:
        w = csv.writer(out_f, lineterminator="\n")
        w.writerow(
            [
                "pro_com",
                "sito_web",
                "email",
                "pec",
                "telefono",
                "codice_fiscale",
                "indirizzo_fisico",
            ]
        )
        for r in sorted(rows_out, key=lambda x: int(x[0])):
            w.writerow(r)
    finally:
        if close_out:
            out_f.close()

    print(
        f"ipa_enti_rows_scanned:{n_total} ipa_comuni_emitted:{n_comuni} phones_mapped:{len(phone_by_comune)}",
        file=sys.stderr,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))

