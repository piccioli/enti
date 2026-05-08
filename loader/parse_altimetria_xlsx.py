#!/usr/bin/env python3
"""
Legge Altimetria Comuni ISTAT (XLSX) e scrive CSV su stdout:
  pro_com,alt_min,alt_max,alt_media,alt_centro
"""
import csv
import sys


def idx_map(header_row):
    m = {}
    for i, h in enumerate(header_row):
        if h is None:
            continue
        name = str(h).strip()
        if name:
            m[name] = i
    return m


def to_int(val):
    if val is None or val == '':
        return None
    try:
        return int(round(float(val)))
    except (TypeError, ValueError):
        return None


def to_decimal(val):
    if val is None or val == '':
        return None
    try:
        return round(float(val), 4)
    except (TypeError, ValueError):
        return None


def main():
    if len(sys.argv) != 2:
        print("Usage: parse_altimetria_xlsx.py FILE.xlsx", file=sys.stderr)
        sys.exit(2)
    path = sys.argv[1]
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("ERROR: python3-openpyxl required", file=sys.stderr)
        sys.exit(1)

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    it = ws.iter_rows(values_only=True)
    next(it, None)  # riga titolo "DATI COMUNALI ..."
    hdr = next(it, None)
    if not hdr:
        print("ERROR: empty sheet", file=sys.stderr)
        sys.exit(1)
    col = idx_map(hdr)
    req = ["PRO_COM", "ALT MIN", "ALT MAX", "MEDIA", "ALT CENTR MUN"]
    for k in req:
        if k not in col:
            print(f"ERROR: colonna '{k}' mancante nell'header ISTAT.", file=sys.stderr)
            sys.exit(1)

    w = csv.writer(sys.stdout)
    w.writerow(["pro_com", "alt_min", "alt_max", "alt_media", "alt_centro"])
    n = 0
    for row in it:
        if not row:
            continue
        try:
            raw_pc = row[col["PRO_COM"]]
        except (IndexError, TypeError):
            continue
        if raw_pc is None or raw_pc == "":
            continue
        try:
            pro_com = int(float(raw_pc))
        except (TypeError, ValueError):
            continue
        alt_min = to_int(row[col["ALT MIN"]])
        alt_max = to_int(row[col["ALT MAX"]])
        alt_media = to_decimal(row[col["MEDIA"]])
        alt_centro = to_int(row[col["ALT CENTR MUN"]])
        w.writerow([pro_com, alt_min if alt_min is not None else "", alt_max if alt_max is not None else "",
                    alt_media if alt_media is not None else "", alt_centro if alt_centro is not None else ""])
        n += 1
    wb.close()
    print(f"wrote_rows={n}", file=sys.stderr)


if __name__ == "__main__":
    main()
